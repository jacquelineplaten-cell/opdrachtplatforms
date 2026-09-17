#!/usr/bin/env python3
"""Leest Salesprofielen_innovatie_groep.xlsx uit naar leesbare tekst per persoon.

profiles.yaml is met de hand verfijnd: de trefwoorden daarin zijn gekozen uit de
antwoorden in de Excel, niet blind overgenomen. Dit script vervangt dat bestand
dus niet. Het helpt je bij bijwerken: het zet de Excel om naar platte tekst per
persoon, zodat je ziet wat er veranderd is en welke trefwoorden je in
profiles.yaml wilt aanpassen.

Gebruik:
    python tools/profielen_uit_excel.py pad/naar/Salesprofielen_innovatie_groep.xlsx
    python tools/profielen_uit_excel.py bestand.xlsx --skelet   # yaml-startpunt
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - alleen bij handmatig gebruik
    sys.exit("openpyxl ontbreekt. Installeer het met: pip install openpyxl")

VRAAG_KOLOM = 1  # kolom A bevat de vragen


def lees(pad: Path, bladnaam: str | None) -> dict[str, dict[str, str]]:
    werkboek = openpyxl.load_workbook(pad, data_only=True)
    blad = werkboek[bladnaam] if bladnaam else werkboek.worksheets[0]

    namen: dict[int, str] = {}
    for kolom in range(VRAAG_KOLOM + 1, blad.max_column + 1):
        waarde = blad.cell(row=1, column=kolom).value
        if waarde and str(waarde).strip():
            namen[kolom] = str(waarde).strip()

    profielen: dict[str, dict[str, str]] = {naam: {} for naam in namen.values()}
    for rij in range(2, blad.max_row + 1):
        vraag = blad.cell(row=rij, column=VRAAG_KOLOM).value
        if not vraag or not str(vraag).strip():
            continue
        vraag = str(vraag).strip()
        for kolom, naam in namen.items():
            waarde = blad.cell(row=rij, column=kolom).value
            tekst = str(waarde).strip() if waarde is not None else ""
            if tekst and not tekst.startswith("#"):  # #VALUE! e.d. overslaan
                profielen[naam][vraag] = tekst
    return profielen


def toon(profielen: dict[str, dict[str, str]]) -> None:
    for naam, antwoorden in profielen.items():
        print(f"\n{'=' * 70}\n{naam}\n{'=' * 70}")
        if not antwoorden:
            print("  (geen profielgegevens ingevuld)")
            continue
        for vraag, antwoord in antwoorden.items():
            print(f"\n-- {vraag}")
            for regel in antwoord.splitlines():
                if regel.strip():
                    print(f"   {regel.strip()}")


def skelet(profielen: dict[str, dict[str, str]]) -> None:
    print("personen:")
    for naam, antwoorden in profielen.items():
        gevuld = bool(antwoorden)
        print(f"  - naam: {naam}")
        print(f"    actief: {'true' if gevuld else 'false'}"
              + ("" if gevuld else "   # geen profielgegevens in de Excel"))
        for veld in ("rollen", "sectoren", "expertise", "uitsluiten"):
            print(f"    {veld}: []")
        print("    boost: {}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("bestand", type=Path)
    parser.add_argument("--blad", help="naam van het werkblad (standaard het eerste)")
    parser.add_argument(
        "--skelet", action="store_true", help="print een leeg yaml-startpunt in plaats van de tekst"
    )
    args = parser.parse_args()

    if not args.bestand.exists():
        sys.exit(f"Bestand niet gevonden: {args.bestand}")

    profielen = lees(args.bestand, args.blad)
    if args.skelet:
        skelet(profielen)
    else:
        toon(profielen)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
